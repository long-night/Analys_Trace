"""Trace 统计分析模块（基于 analys_trace_v4.py 重构）"""
import csv
import os
import sys
from collections import defaultdict
from typing import Dict, List, Optional

from op_testgen.parser.trace_parser import OpInfo, HierarchicalOpInfo, TraceParser
from op_testgen.analyzer.hierarchy_analyzer import HierarchyAnalyzer


class ShapeStats:
    """单个 Shape 维度的统计信息"""

    def __init__(self, shape: tuple, strides: Optional[tuple] = None):
        self.shape = shape
        self.strides = strides
        self.call_count = 0
        self.total_duration_us = 0.0
        self.durations_us: List[float] = []
        self.input_types_set: set = set()
        self.concrete_inputs_set: set = set()

    def add_call(self, duration_us: float, input_types: Optional[List[str]] = None,
                 concrete_inputs: Optional[List] = None):
        self.call_count += 1
        self.total_duration_us += duration_us
        self.durations_us.append(duration_us)
        if input_types:
            self.input_types_set.add(str(input_types))
        if concrete_inputs:
            self.concrete_inputs_set.add(str(concrete_inputs))

    @property
    def avg_duration_ms(self) -> float:
        return (self.total_duration_us / self.call_count / 1000.0) if self.call_count > 0 else 0.0

    @property
    def min_duration_ms(self) -> float:
        return min(self.durations_us) / 1000.0 if self.durations_us else 0.0

    @property
    def max_duration_ms(self) -> float:
        return max(self.durations_us) / 1000.0 if self.durations_us else 0.0

    @property
    def total_duration_ms(self) -> float:
        return self.total_duration_us / 1000.0


class OperatorStats:
    """单个算子的聚合统计信息"""

    def __init__(self, name: str, depth: int = 0, root_name: str = ""):
        self.name = name
        self.depth = depth
        self.root_name = root_name
        self.call_count = 0
        self.total_duration_us = 0.0
        self.shape_stats: Dict[str, ShapeStats] = {}

    def add_op_info(self, op_info: OpInfo):
        """从 OpInfo 添加统计"""
        self.call_count += 1
        self.total_duration_us += op_info.duration_us

        if op_info.input_dims:
            shape_tuple = tuple(tuple(d) for d in op_info.input_dims)
            strides_tuple = tuple(tuple(s) for s in op_info.input_strides) if op_info.input_strides else None
            key = str(list(shape_tuple))
            if strides_tuple:
                key += " | " + str(list(strides_tuple))

            if key not in self.shape_stats:
                self.shape_stats[key] = ShapeStats(shape_tuple, strides_tuple)

            self.shape_stats[key].add_call(
                op_info.duration_us,
                op_info.input_types,
                op_info.concrete_inputs
            )

    @property
    def total_duration_ms(self) -> float:
        return self.total_duration_us / 1000.0

    @property
    def avg_duration_ms(self) -> float:
        return (self.total_duration_us / self.call_count / 1000.0) if self.call_count > 0 else 0.0

    def get_all_shapes_str(self) -> str:
        if not self.shape_stats:
            return "N/A"
        shapes = {str(list(s.shape)) for s in self.shape_stats.values()}
        return " | ".join(sorted(shapes))

    def get_all_strides_str(self) -> str:
        if not self.shape_stats:
            return "N/A"
        strides = {str(list(s.strides)) for s in self.shape_stats.values() if s.strides}
        return " | ".join(sorted(strides)) if strides else "N/A"

    def get_all_input_types_str(self) -> str:
        if not self.shape_stats:
            return "N/A"
        types = set()
        for s in self.shape_stats.values():
            types.update(s.input_types_set)
        return " | ".join(sorted(types)) if types else "N/A"

    def get_all_concrete_inputs_str(self) -> str:
        if not self.shape_stats:
            return "N/A"
        concretes = set()
        for s in self.shape_stats.values():
            concretes.update(s.concrete_inputs_set)
        return " | ".join(sorted(concretes)) if concretes else "N/A"


class TraceAnalyzer:
    """Trace 统计分析器"""

    def __init__(self, op_infos: List[OpInfo], use_hierarchical: bool = False):
        self.op_infos = op_infos
        self.operators: Dict[str, OperatorStats] = {}
        self.use_hierarchical = use_hierarchical
        self._hierarchy: Optional[List[HierarchicalOpInfo]] = None
        self._hierarchy_analyzer: Optional[HierarchyAnalyzer] = None
        self._analyze()

    def _analyze(self):
        """聚合统计"""
        for op_info in self.op_infos:
            if self.use_hierarchical and isinstance(op_info, HierarchicalOpInfo):
                name = op_info.hierarchical_name
                depth = op_info.depth
                root_name = op_info.get_call_chain()[0] if op_info.get_call_chain() else name
            else:
                name = op_info.name
                depth = 0
                root_name = name

            if name not in self.operators:
                self.operators[name] = OperatorStats(name, depth, root_name)
            self.operators[name].add_op_info(op_info)

    def print_summary(self):
        """打印摘要信息"""
        print("\n" + "=" * 60)
        print("分析摘要".center(60))
        print("=" * 60)

        total_time_all = sum(op.total_duration_us for op in self.operators.values())

        # 根算子统计（层级模式下）
        if self.use_hierarchical:
            root_ops = {name: op for name, op in self.operators.items() if op.depth == 0}
            if root_ops:
                print("\n【根算子统计】")
                print(f"  根算子数: {len(root_ops)}")
                print(f"  根算子总时间: {sum(op.total_duration_us for op in root_ops.values()) / 1000.0:.2f} ms")
                print(f"  总算子（含层级）: {len(self.operators)}")

        # 最耗时的算子 TOP 10
        print("\n【最耗时的算子 TOP 10】")
        top_time = sorted(self.operators.values(), key=lambda x: x.total_duration_us, reverse=True)[:10]
        for i, op in enumerate(top_time, 1):
            pct = (op.total_duration_us / total_time_all * 100) if total_time_all > 0 else 0
            depth_tag = f" [深度 {op.depth}]" if self.use_hierarchical else ""
            print(f"  {i:2d}. {op.name}{depth_tag}")
            print(f"      总时间: {op.total_duration_ms:.2f} ms ({pct:.1f}%), "
                  f"调用: {op.call_count} 次, 平均: {op.avg_duration_ms:.4f} ms")

        # 调用最频繁的算子 TOP 5
        print("\n【调用最频繁的算子 TOP 5】")
        top_call = sorted(self.operators.values(), key=lambda x: x.call_count, reverse=True)[:5]
        for i, op in enumerate(top_call, 1):
            print(f"  {i}. {op.name}: {op.call_count} 次, 总时间: {op.total_duration_ms:.2f} ms")

        # Shape 多样性
        print("\n【Shape 多样性最高的算子 TOP 5】")
        top_diversity = sorted(self.operators.values(), key=lambda x: len(x.shape_stats), reverse=True)[:5]
        for i, op in enumerate(top_diversity, 1):
            print(f"  {i}. {op.name}: {len(op.shape_stats)} 种不同的 shape")

        # 通信算子统计
        comm_ops = {n: o for n, o in self.operators.items()
                    if any(n.startswith(p) for p in ("c10d::", "nccl:", "gloo:", "mpi:"))}
        if comm_ops:
            print("\n【通信算子统计】")
            comm_time = sum(o.total_duration_us for o in comm_ops.values())
            comm_pct = (comm_time / total_time_all * 100) if total_time_all > 0 else 0
            print(f"  通信算子数: {len(comm_ops)}")
            print(f"  通信总时间: {comm_time / 1000.0:.2f} ms ({comm_pct:.1f}%)")
            print(f"  通信调用次数: {sum(o.call_count for o in comm_ops.values())}")
            print("\n  通信算子 TOP 5:")
            for i, op in enumerate(sorted(comm_ops.values(), key=lambda x: x.total_duration_us, reverse=True)[:5], 1):
                print(f"    {i}. {op.name}: {op.total_duration_ms:.2f} ms, 调用: {op.call_count} 次")

        # 总体统计
        print(f"\n【总体统计】")
        print(f"  总算子数: {len(self.operators)}")
        print(f"  总执行时间: {total_time_all / 1000.0:.2f} ms")
        print(f"  总调用次数: {sum(op.call_count for op in self.operators.values())}")
        print("=" * 60)

    def export_to_csv(self, operators_file: str, shapes_file: str):
        """导出到 CSV"""
        sorted_ops = sorted(self.operators.values(), key=lambda x: x.total_duration_us, reverse=True)

        with open(operators_file, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=[
                "算子名称", "深度", "根算子", "输入Shapes", "输入Strides", "输入数据类型", "Concrete Inputs",
                "调用次数", "总执行时间(ms)", "平均执行时间(ms)"
            ])
            writer.writeheader()
            for op in sorted_ops:
                writer.writerow({
                    "算子名称": op.name,
                    "深度": op.depth,
                    "根算子": op.root_name,
                    "输入Shapes": op.get_all_shapes_str(),
                    "输入Strides": op.get_all_strides_str(),
                    "输入数据类型": op.get_all_input_types_str(),
                    "Concrete Inputs": op.get_all_concrete_inputs_str(),
                    "调用次数": op.call_count,
                    "总执行时间(ms)": round(op.total_duration_ms, 4),
                    "平均执行时间(ms)": round(op.avg_duration_ms, 4),
                })

        all_shape_stats = []
        for op in self.operators.values():
            for stats in op.shape_stats.values():
                all_shape_stats.append({
                    "op_name": op.name,
                    "depth": op.depth,
                    "root_name": op.root_name,
                    "shape": str(list(stats.shape)),
                    "strides": str(list(stats.strides)) if stats.strides else "N/A",
                    "input_types": " | ".join(sorted(stats.input_types_set)) if stats.input_types_set else "N/A",
                    "concrete_inputs": " | ".join(sorted(stats.concrete_inputs_set)) if stats.concrete_inputs_set else "N/A",
                    "call_count": stats.call_count,
                    "total_duration_ms": stats.total_duration_ms,
                    "avg_duration_ms": stats.avg_duration_ms,
                    "min_duration_ms": stats.min_duration_ms,
                    "max_duration_ms": stats.max_duration_ms,
                })
        all_shape_stats.sort(key=lambda x: (x["op_name"], -float(x["total_duration_ms"])))

        with open(shapes_file, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=[
                "算子名称", "深度", "根算子", "输入Shapes", "输入Strides", "输入数据类型", "Concrete Inputs",
                "调用次数", "总执行时间(ms)", "平均执行时间(ms)",
                "最小执行时间(ms)", "最大执行时间(ms)"
            ])
            writer.writeheader()
            for s in all_shape_stats:
                writer.writerow({
                    "算子名称": s["op_name"],
                    "深度": s["depth"],
                    "根算子": s["root_name"],
                    "输入Shapes": s["shape"],
                    "输入Strides": s["strides"],
                    "输入数据类型": s["input_types"],
                    "Concrete Inputs": s["concrete_inputs"],
                    "调用次数": s["call_count"],
                    "总执行时间(ms)": round(s["total_duration_ms"], 4),
                    "平均执行时间(ms)": round(s["avg_duration_ms"], 4),
                    "最小执行时间(ms)": round(s["min_duration_ms"], 4),
                    "最大执行时间(ms)": round(s["max_duration_ms"], 4),
                })

    def export_to_excel(self, output_file: str) -> bool:
        """导出到 Excel"""
        try:
            import openpyxl
            from openpyxl.styles import Font, Alignment, PatternFill
            from openpyxl.utils import get_column_letter
        except ImportError:
            print("错误: openpyxl 未安装，无法导出 Excel")
            print("请运行: pip install openpyxl")
            return False

        wb = openpyxl.Workbook()
        if wb.active is not None:
            wb.remove(wb.active)

        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        header_align = Alignment(horizontal="center", vertical="center")

        # Sheet 1: 算子总表
        ws1 = wb.create_sheet("算子总表")
        headers1 = ["算子名称", "深度", "根算子", "输入Shapes", "输入Strides", "输入数据类型", "Concrete Inputs",
                    "调用次数", "总执行时间(ms)", "平均执行时间(ms)"]
        ws1.append(headers1)
        for col_num, _ in enumerate(headers1, 1):
            cell = ws1.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_align

        sorted_ops = sorted(self.operators.values(), key=lambda x: x.total_duration_us, reverse=True)
        for op in sorted_ops:
            ws1.append([
                op.name, op.depth, op.root_name,
                op.get_all_shapes_str(), op.get_all_strides_str(),
                op.get_all_input_types_str(), op.get_all_concrete_inputs_str(),
                op.call_count, round(op.total_duration_ms, 4), round(op.avg_duration_ms, 4)
            ])

        # Sheet 2: Shape 统计表
        ws2 = wb.create_sheet("Shape统计表")
        headers2 = ["算子名称", "深度", "根算子", "输入Shapes", "输入Strides", "输入数据类型", "Concrete Inputs",
                    "调用次数", "总执行时间(ms)", "平均执行时间(ms)",
                    "最小执行时间(ms)", "最大执行时间(ms)"]
        ws2.append(headers2)
        for col_num, _ in enumerate(headers2, 1):
            cell = ws2.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_align

        all_shape_stats = []
        for op in self.operators.values():
            for stats in op.shape_stats.values():
                all_shape_stats.append({
                    "op_name": op.name,
                    "depth": op.depth,
                    "root_name": op.root_name,
                    "shape": str(list(stats.shape)),
                    "strides": str(list(stats.strides)) if stats.strides else "N/A",
                    "input_types": " | ".join(sorted(stats.input_types_set)) if stats.input_types_set else "N/A",
                    "concrete_inputs": " | ".join(sorted(stats.concrete_inputs_set)) if stats.concrete_inputs_set else "N/A",
                    "call_count": stats.call_count,
                    "total_duration_ms": stats.total_duration_ms,
                    "avg_duration_ms": stats.avg_duration_ms,
                    "min_duration_ms": stats.min_duration_ms,
                    "max_duration_ms": stats.max_duration_ms,
                })
        all_shape_stats.sort(key=lambda x: (x["op_name"], -float(x["total_duration_ms"])))

        for s in all_shape_stats:
            ws2.append([
                s["op_name"], s["depth"], s["root_name"],
                s["shape"], s["strides"], s["input_types"], s["concrete_inputs"],
                s["call_count"], round(float(s["total_duration_ms"]), 4), round(float(s["avg_duration_ms"]), 4),
                round(float(s["min_duration_ms"]), 4), round(float(s["max_duration_ms"]), 4)
            ])

        # 自动调整列宽
        for ws in [ws1, ws2]:
            for col in ws.columns:
                max_length = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    try:
                        if cell.value:
                            max_length = max(max_length, len(str(cell.value)))
                    except Exception:
                        pass
                ws.column_dimensions[col_letter].width = min(max_length + 2, 80)

        wb.save(output_file)
        return True

    def _ensure_hierarchy(self):
        if self._hierarchy is None:
            parser = TraceParser("")
            self._hierarchy = parser._build_hierarchy(self.op_infos)
            self._hierarchy_analyzer = HierarchyAnalyzer(self._hierarchy)

    def print_hierarchical_summary(self):
        self._ensure_hierarchy()
        assert self._hierarchy_analyzer is not None

        print("\n" + "=" * 60)
        print("层级分析摘要".center(60))
        print("=" * 60)

        self_time_stats = self._hierarchy_analyzer.get_self_time_stats()
        total_self_time = sum(s.total_duration_us for s in self_time_stats.values())

        print("\n【自耗时最高的算子 TOP 10】（排除子算子影响）")
        top_self = sorted(self_time_stats.values(), key=lambda x: x.total_duration_us, reverse=True)[:10]
        for i, op in enumerate(top_self, 1):
            pct = (op.total_duration_us / total_self_time * 100) if total_self_time > 0 else 0
            print(f"  {i:2d}. {op.name}")
            print(f"      自耗时: {op.total_duration_ms:.2f} ms ({pct:.1f}%), 调用: {op.call_count} 次")

        chains = self._hierarchy_analyzer.get_call_chains()
        if chains:
            print("\n【高频调用链 TOP 5】")
            for i, chain in enumerate(chains[:5], 1):
                chain_str = " -> ".join(chain.chain)
                print(f"  {i}. {chain_str} (出现 {chain.occurrence_count} 次, 平均耗时: {chain.avg_duration_ms:.2f} ms)")

        patterns = self._hierarchy_analyzer.find_recursive_patterns()
        if patterns:
            print("\n【检测到的循环调用模式】")
            for i, pattern in enumerate(patterns[:5], 1):
                cycle_str = " -> ".join(pattern.cycle)
                print(f"  {i}. {cycle_str} (出现 {pattern.occurrence_count} 次, 最大深度: {pattern.depth})")

        depth_dist = self._hierarchy_analyzer.get_depth_distribution()
        if depth_dist:
            print("\n【调用深度分布】")
            for depth, count in depth_dist.items():
                label = "根" if depth == 0 else f"深度 {depth}"
                print(f"  {label}: {count} 个算子")

        print("=" * 60)

    def export_hierarchy_to_csv(self, output_file: str):
        self._ensure_hierarchy()
        assert self._hierarchy_analyzer is not None

        with open(output_file, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=[
                "算子名称", "调用深度", "是否为根", "自耗时(ms)", "总耗时(ms)",
                "子算子数", "调用次数"
            ])
            writer.writeheader()

            for node in self._hierarchy_analyzer.all_nodes:
                writer.writerow({
                    "算子名称": node.name,
                    "调用深度": node.depth,
                    "是否为根": "是" if node.is_root else "否",
                    "自耗时(ms)": round(node.self_duration_us / 1000.0, 4),
                    "总耗时(ms)": round(node.duration_us / 1000.0, 4),
                    "子算子数": len(node.children),
                    "调用次数": 1,
                })

    def export_hierarchy_to_excel(self, output_file: str) -> bool:
        self._ensure_hierarchy()
        assert self._hierarchy_analyzer is not None

        try:
            import openpyxl
            from openpyxl.styles import Font, Alignment, PatternFill
            from openpyxl.utils import get_column_letter
        except ImportError:
            print("错误: openpyxl 未安装，无法导出 Excel")
            return False

        wb = openpyxl.load_workbook(output_file)

        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        header_align = Alignment(horizontal="center", vertical="center")

        ws_self = wb.create_sheet("自耗时排名")
        headers = ["算子名称", "调用深度", "是否为根", "自耗时(ms)", "总耗时(ms)", "子算子数"]
        ws_self.append(headers)
        for col_num, _ in enumerate(headers, 1):
            cell = ws_self.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_align

        for node in self._hierarchy_analyzer.all_nodes:
            ws_self.append([
                node.name, node.depth, "是" if node.is_root else "否",
                round(node.self_duration_us / 1000.0, 4),
                round(node.duration_us / 1000.0, 4),
                len(node.children),
            ])

        for col in ws_self.columns:
            max_length = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                try:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except Exception:
                    pass
            ws_self.column_dimensions[col_letter].width = min(max_length + 2, 80)

        wb.save(output_file)
        return True
