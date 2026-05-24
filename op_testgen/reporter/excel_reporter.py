"""Excel 报告生成器"""
from typing import List

from op_testgen.correctness.test_runner import CorrectnessResult
from op_testgen.perf.benchmark import PerfResult
from op_testgen.reporter.base import BaseReporter


class ExcelReporter(BaseReporter):
    """Excel 报告生成器（需要 openpyxl）"""

    def generate(self, correctness_results: List[CorrectnessResult],
                 perf_results: List[PerfResult], output_path: str) -> None:
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
        except ImportError:
            raise ImportError("Excel 报告需要 openpyxl，请安装: pip install openpyxl")

        wb = openpyxl.Workbook()

        ws1 = wb.active
        ws1.title = "正确性测试"
        headers1 = ["算子名称", "后端", "通过状态", "最大绝对误差", "最大相对误差", "平均绝对误差", "错误信息"]
        ws1.append(headers1)

        for r in correctness_results:
            ws1.append([
                r.op_name,
                r.backend,
                "通过" if r.passed else "失败",
                r.max_abs_err,
                r.max_rel_err,
                r.avg_abs_err,
                r.error_message or "",
            ])

        ws2 = wb.create_sheet("性能测试")
        headers2 = ["算子名称", "分类", "平均耗时(ms)", "GFLOPS", "带宽(GB/s)", "加速比"]
        ws2.append(headers2)

        for r in perf_results:
            ws2.append([
                r.op_name,
                r.category,
                r.avg_time_ms,
                r.flops,
                r.bandwidth_gbps,
                r.speedup,
            ])

        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        for ws in [ws1, ws2]:
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center")

        wb.save(output_path)
